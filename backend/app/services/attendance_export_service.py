"""
attendance_export_service.py - Enterprise Attendance & Truancy Audit Export Engine.
Provides filtered, preset-driven exports for:
1. Class Attendance Register Matrix (Student vs Dates)
2. Chronic Truancy & Absenteeism Alert Roster (<75% attendance)
3. Subject / Period Lesson Cut Truancy Audit

Supports both styled .xlsx (openpyxl) and universal .csv with strict school-level isolation.
100% Offline-First.
"""

import io
import csv
import re
from datetime import datetime, date, timedelta
from typing import Any, Dict, List, Optional, Tuple
from sqlalchemy.orm import Session, joinedload
from sqlalchemy import or_, and_, func, desc

import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter

from ..models import Attendance, Student, ClassSection, Subject, School, User
from .import_export_service import sanitize_csv_cell


def build_filtered_attendance_query(
    db: Session,
    school_id: Optional[int],
    filters: Optional[Dict[str, Any]] = None
):
    """
    Builds a tenant-isolated query for Attendance records.
    """
    filters = filters or {}
    query = db.query(Attendance).join(Student, Attendance.student_id == Student.id).options(
        joinedload(Attendance.student).joinedload(Student.class_section),
        joinedload(Attendance.subject)
    )

    if school_id is not None:
        query = query.filter(Student.school_id == school_id)

    # 1. Class filter
    class_id = filters.get("class_id")
    if class_id is not None and str(class_id).strip() != "":
        try:
            query = query.filter(Student.class_section_id == int(class_id))
        except (ValueError, TypeError):
            pass

    # 2. Attendance Type (daily vs period)
    att_type = filters.get("attendance_type")
    if att_type and str(att_type).strip().lower() not in ("all", ""):
        query = query.filter(Attendance.attendance_type == str(att_type).strip().lower())

    # 3. Subject filter
    subject_id = filters.get("subject_id")
    if subject_id is not None and str(subject_id).strip() != "":
        try:
            query = query.filter(Attendance.subject_id == int(subject_id))
        except (ValueError, TypeError):
            pass

    # 4. Date Range
    start_date = filters.get("start_date")
    if start_date and str(start_date).strip() != "":
        try:
            dt_start = datetime.strptime(str(start_date).strip()[:10], "%Y-%m-%d")
            query = query.filter(Attendance.date >= dt_start)
        except ValueError:
            pass

    end_date = filters.get("end_date")
    if end_date and str(end_date).strip() != "":
        try:
            dt_end = datetime.strptime(str(end_date).strip()[:10] + " 23:59:59", "%Y-%m-%d %H:%M:%S")
            query = query.filter(Attendance.date <= dt_end)
        except ValueError:
            pass

    return query


def generate_attendance_export_dataset(
    db: Session,
    school_id: Optional[int],
    preset_key: str = "register_matrix",
    file_format: str = "xlsx",
    filters: Optional[Dict[str, Any]] = None
) -> Tuple[bytes, str, str]:
    """
    Generates preset-driven attendance exports:
    - register_matrix
    - truancy_alert
    - subject_cuts
    """
    filters = filters or {}
    school = db.query(School).filter(School.id == school_id).first() if school_id else None
    school_slug = re.sub(r"[^a-zA-Z0-9_-]", "_", school.name if school else "Institutional")
    timestamp = datetime.now().strftime("%Y%m%d_%H%M")

    # ── Preset 1: Subject / Lesson Cut Audit ──
    if preset_key == "subject_cuts":
        title = "Subject_Period_Absence_Audit"
        columns = [
            ("Subject", "subject_name"),
            ("Class Section", "class_name"),
            ("Period / Lesson", "period_label"),
            ("Date", "date_str"),
            ("Student Code", "student_code"),
            ("Student Name", "full_name"),
            ("Status", "status")
        ]

        query = build_filtered_attendance_query(db, school_id, filters)
        query = query.filter(
            Attendance.attendance_type == "period",
            Attendance.status.in_(["Absent", "Late"])
        )
        records = query.order_by(desc(Attendance.date), Student.full_name.asc()).all()

        rows_data = []
        for att in records:
            st = att.student
            sub = att.subject
            rows_data.append({
                "subject_name": sub.name if sub else "Unspecified Subject",
                "class_name": st.class_section.name if (st and st.class_section) else "",
                "period_label": att.period_label or "Period Lesson",
                "date_str": att.date.strftime("%Y-%m-%d") if att.date else "",
                "student_code": st.student_code if st else "",
                "full_name": st.full_name if st else "",
                "status": att.status or "Absent"
            })

        return _render_attendance_dataset(rows_data, columns, f"{title}_{school_slug}_{timestamp}", file_format)

    # ── Preset 2: Chronic Truancy Alert Roster ──
    elif preset_key == "truancy_alert":
        title = "Chronic_Truancy_Absenteeism_Alert"
        columns = [
            ("Student Code", "student_code"),
            ("Full Name", "full_name"),
            ("Gender", "gender"),
            ("Class Section", "class_name"),
            ("Total Sessions", "total_days"),
            ("Sessions Present", "present_days"),
            ("Sessions Absent", "absent_days"),
            ("Attendance Rate (%)", "att_rate"),
            ("Guardian Name", "guardian_name"),
            ("Guardian Phone", "phone"),
            ("Residential Status", "residential_status")
        ]

        class_id = filters.get("class_id")
        threshold = 75.0
        thresh_val = filters.get("threshold")
        if thresh_val is not None:
            try:
                threshold = float(thresh_val)
            except (ValueError, TypeError):
                threshold = 75.0

        st_query = db.query(Student).filter(Student.is_active == True).options(
            joinedload(Student.class_section),
            joinedload(Student.attendance)
        )
        if school_id is not None:
            st_query = st_query.filter(Student.school_id == school_id)
        if class_id:
            try:
                st_query = st_query.filter(Student.class_section_id == int(class_id))
            except (ValueError, TypeError):
                pass

        students = st_query.order_by(Student.full_name.asc()).all()
        rows_data = []
        for s in students:
            # Calculate daily attendance
            daily_att = [a for a in s.attendance if a.attendance_type == "daily"]
            tot = len(daily_att)
            if tot == 0:
                continue
            pres = sum(1 for a in daily_att if a.status in ("Present", "Late"))
            absent = tot - pres
            rate = round((pres / tot) * 100.0, 1)

            if rate < threshold:
                rows_data.append({
                    "student_code": s.student_code or "",
                    "full_name": s.full_name or "",
                    "gender": s.gender or "",
                    "class_name": s.class_section.name if s.class_section else "",
                    "total_days": tot,
                    "present_days": pres,
                    "absent_days": absent,
                    "att_rate": f"{rate}%",
                    "guardian_name": s.guardian_name or "",
                    "phone": s.phone or "",
                    "residential_status": s.residential_status or ""
                })

        rows_data.sort(key=lambda x: float(x["att_rate"].replace("%", "")))
        return _render_attendance_dataset(rows_data, columns, f"{title}_{school_slug}_{timestamp}", file_format)

    # ── Preset 3 (Default): Class Attendance Register Matrix ──
    else:
        title = "Class_Attendance_Register_Matrix"
        class_id = filters.get("class_id")

        st_query = db.query(Student).filter(Student.is_active == True).options(
            joinedload(Student.class_section)
        )
        if school_id is not None:
            st_query = st_query.filter(Student.school_id == school_id)
        if class_id:
            try:
                st_query = st_query.filter(Student.class_section_id == int(class_id))
            except (ValueError, TypeError):
                pass

        students = st_query.order_by(Student.full_name.asc()).all()
        if not students:
            return _render_attendance_dataset([], [("Notice", "notice")], f"{title}_{school_slug}_{timestamp}", file_format)

        st_ids = [s.id for s in students]

        # Fetch attendance records
        att_query = db.query(Attendance).filter(
            Attendance.student_id.in_(st_ids),
            Attendance.attendance_type == "daily"
        )
        start_date = filters.get("start_date")
        if start_date:
            try:
                att_query = att_query.filter(Attendance.date >= datetime.strptime(str(start_date)[:10], "%Y-%m-%d"))
            except ValueError:
                pass
        end_date = filters.get("end_date")
        if end_date:
            try:
                att_query = att_query.filter(Attendance.date <= datetime.strptime(str(end_date)[:10] + " 23:59:59", "%Y-%m-%d %H:%M:%S"))
            except ValueError:
                pass

        att_records = att_query.all()

        # Collect unique sorted dates
        date_set = set()
        st_date_status: Dict[int, Dict[str, str]] = {}
        for a in att_records:
            d_str = a.date.strftime("%Y-%m-%d") if a.date else ""
            if d_str:
                date_set.add(d_str)
                # P, A, L, E code
                code = "P" if a.status == "Present" else ("A" if a.status == "Absent" else ("L" if a.status == "Late" else "E"))
                st_date_status.setdefault(a.student_id, {})[d_str] = code

        sorted_dates = sorted(list(date_set))

        # Dynamic Columns
        columns = [
            ("Student Code", "student_code"),
            ("Full Name", "full_name"),
            ("Gender", "gender"),
            ("Class Section", "class_name")
        ]
        for d in sorted_dates:
            columns.append((d, f"d_{d}"))

        columns.append(("Present", "present_days"))
        columns.append(("Absent", "absent_days"))
        columns.append(("Rate (%)", "att_rate"))

        rows_data = []
        for s in students:
            row_dict = {
                "student_code": s.student_code or "",
                "full_name": s.full_name or "",
                "gender": s.gender or "",
                "class_name": s.class_section.name if s.class_section else ""
            }
            pres = 0
            absent = 0
            s_map = st_date_status.get(s.id, {})
            for d in sorted_dates:
                code = s_map.get(d, "-")
                row_dict[f"d_{d}"] = code
                if code in ("P", "L"):
                    pres += 1
                elif code == "A":
                    absent += 1

            total_days = pres + absent
            rate = round((pres / total_days * 100.0), 1) if total_days > 0 else 100.0
            row_dict["present_days"] = pres
            row_dict["absent_days"] = absent
            row_dict["att_rate"] = f"{rate}%"
            rows_data.append(row_dict)

        return _render_attendance_dataset(rows_data, columns, f"{title}_{school_slug}_{timestamp}", file_format)


def _render_attendance_dataset(
    rows_data: List[Dict[str, Any]],
    columns: List[Tuple[str, str]],
    filename_base: str,
    file_format: str
) -> Tuple[bytes, str, str]:
    """Helper to render styled Excel or CSV for attendance rosters."""
    if file_format.lower() == "xlsx":
        wb = openpyxl.Workbook()
        ws = wb.active
        ws.title = "Attendance_Audit"
        ws.views.sheetView[0].showGridLines = True

        header_font = Font(name="Segoe UI", size=10, bold=True, color="FFFFFF")
        header_fill = PatternFill(start_color="1E293B", end_color="1E293B", fill_type="solid")
        cell_font = Font(name="Segoe UI", size=10)
        center_align = Alignment(horizontal="center", vertical="center")
        left_align = Alignment(horizontal="left", vertical="center")
        right_align = Alignment(horizontal="right", vertical="center")
        thin_border = Border(
            left=Side(style="thin", color="CBD5E1"),
            right=Side(style="thin", color="CBD5E1"),
            top=Side(style="thin", color="CBD5E1"),
            bottom=Side(style="thin", color="CBD5E1")
        )
        zebra_fill = PatternFill(start_color="F8FAFC", end_color="F8FAFC", fill_type="solid")

        # Header
        ws.row_dimensions[1].height = 26
        for col_idx, (col_label, col_key) in enumerate(columns, start=1):
            cell = ws.cell(row=1, column=col_idx, value=col_label)
            cell.font = header_font
            cell.fill = header_fill
            cell.alignment = center_align
            cell.border = thin_border

        # Data rows
        for row_idx, r in enumerate(rows_data, start=2):
            ws.row_dimensions[row_idx].height = 20
            is_even = (row_idx % 2 == 0)

            for col_idx, (col_label, col_key) in enumerate(columns, start=1):
                val = r.get(col_key, "")
                cell = ws.cell(row=row_idx, column=col_idx, value=val)
                cell.font = cell_font
                cell.border = thin_border
                if is_even:
                    cell.fill = zebra_fill

                if isinstance(val, (int, float)):
                    cell.alignment = right_align
                elif col_key in ("student_code", "gender", "status", "att_rate", "date_str") or col_key.startswith("d_"):
                    cell.alignment = center_align
                    if val == "A":
                        cell.font = Font(name="Segoe UI", size=10, bold=True, color="DC2626")
                    elif val in ("P", "L"):
                        cell.font = Font(name="Segoe UI", size=10, bold=True, color="059669")
                else:
                    cell.alignment = left_align

        # Auto-fit columns
        for col in ws.columns:
            max_len = 0
            col_letter = get_column_letter(col[0].column)
            for cell in col:
                val_str = str(cell.value or "")
                if len(val_str) > max_len:
                    max_len = len(val_str)
            ws.column_dimensions[col_letter].width = max(max_len + 3, 7)

        output = io.BytesIO()
        wb.save(output)
        output.seek(0)
        return output.getvalue(), f"{filename_base}.xlsx", "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"

    else:
        output = io.StringIO()
        writer = csv.writer(output, quoting=csv.QUOTE_MINIMAL)
        writer.writerow([c[0] for c in columns])

        for r in rows_data:
            row_vals = [sanitize_csv_cell(r.get(c[1], "")) for c in columns]
            writer.writerow(row_vals)

        csv_bytes = output.getvalue().encode("utf-8-sig")
        return csv_bytes, f"{filename_base}.csv", "text/csv; charset=utf-8"
