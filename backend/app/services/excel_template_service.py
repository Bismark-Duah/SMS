"""
Dynamic Institutional Excel Template Generation & Ingestion Service
Generates school-tailored .xlsx templates with native Data Validation dropdowns
populated dynamically from live school database records:
- Accredited Subjects (Super-Admin mapped + School-Admin custom subjects)
- School Departments
- School Houses
- School Class Sections
- Standard Genders & Institutional Roles

Also parses both .xlsx and .csv uploads with zero external network dependencies.
"""

import io
import os
import re
from typing import Any, Dict, List, Optional, Tuple
from fastapi import HTTPException, UploadFile
from sqlalchemy.orm import Session
import openpyxl
from openpyxl.worksheet.datavalidation import DataValidation
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter

from ..models import School, Subject, Department, House, ClassSection
from .import_export_service import decode_csv_bytes, sanitize_filename


def get_school_onboarding_lookups(db: Session, target_sch_id: Optional[int]) -> Dict[str, List[str]]:
    """
    Extracts live, school-scoped reference lists for template dropdowns.
    Unions Super-Admin accredited curriculum with School-Admin registered subjects.
    """
    school = db.query(School).filter(School.id == target_sch_id).first() if target_sch_id else None

    # 1. Subjects: Super Admin accredited (school.active_subjects) + School Admin custom (Subject.school_id == target_sch_id)
    subject_names = set()
    if school and school.active_subjects:
        for s in school.active_subjects:
            if s.is_active or s.is_active is None:
                if s.name and s.name.strip():
                    subject_names.add(s.name.strip())

    if target_sch_id:
        local_subjs = db.query(Subject).filter(
            (Subject.school_id == target_sch_id) & 
            ((Subject.is_active == True) | (Subject.is_active.is_(None)))
        ).all()
        for s in local_subjs:
            if s.name and s.name.strip():
                subject_names.add(s.name.strip())

    # Fallback to general active subjects if no school-specific subjects found
    if not subject_names:
        for s in db.query(Subject).filter((Subject.is_active == True) | (Subject.is_active.is_(None))).all():
            if s.name and s.name.strip():
                subject_names.add(s.name.strip())

    subjects_list = sorted(list(subject_names))

    # 2. Departments
    dept_query = db.query(Department)
    if target_sch_id:
        depts = dept_query.filter(Department.school_id == target_sch_id).order_by(Department.name.asc()).all()
        if not depts:
            depts = dept_query.filter(Department.school_id.is_(None)).order_by(Department.name.asc()).all()
    else:
        depts = dept_query.order_by(Department.name.asc()).all()
    departments_list = [d.name.strip() for d in depts if d.name and d.name.strip()]

    # 3. Houses
    house_query = db.query(House)
    if target_sch_id:
        houses = house_query.filter(House.school_id == target_sch_id).order_by(House.name.asc()).all()
        if not houses:
            houses = house_query.filter(House.school_id.is_(None)).order_by(House.name.asc()).all()
    else:
        houses = house_query.order_by(House.name.asc()).all()
    houses_list = [h.name.strip() for h in houses if h.name and h.name.strip()]

    # 4. Class Sections
    class_query = db.query(ClassSection)
    if target_sch_id:
        classes = class_query.filter(ClassSection.school_id == target_sch_id).order_by(ClassSection.name.asc()).all()
    else:
        classes = class_query.order_by(ClassSection.name.asc()).all()
    classes_list = [c.name.strip() for c in classes if c.name and c.name.strip()]

    # 5. Genders
    genders_list = ["Male", "Female"]

    # 6. Roles
    roles_list = [
        "teacher",
        "teacher|hod",
        "teacher|form_master",
        "teacher|senior_house_master",
        "teacher|house_master",
        "teacher|assistant_house_master",
        "assistant_headmaster_academic",
        "assistant_headmaster_domestic",
        "assistant_headmaster_admin",
        "school_administrator",
        "ict_coordinator",
        "bursar",
        "storekeeper",
        "security_officer",
        "secretary",
        "parent"
    ]

    return {
        "school_name": school.name if school else "Institutional",
        "subjects": subjects_list,
        "departments": departments_list,
        "houses": houses_list,
        "classes": classes_list,
        "genders": genders_list,
        "roles": roles_list
    }


def generate_staff_onboarding_excel(db: Session, target_sch_id: Optional[int]) -> Tuple[bytes, str]:
    """
    Generates a production-grade .xlsx workbook with two sheets:
    - Staff_Onboarding: The user-facing grid with styled headers and native dropdown validations.
    - Lookups: The school-scoped reference lists (hidden to prevent user tampering).
    """
    lookups = get_school_onboarding_lookups(db, target_sch_id)

    wb = openpyxl.Workbook()
    
    # ── Sheet 1: Staff Onboarding ──
    ws_main = wb.active
    ws_main.title = "Staff_Onboarding"
    ws_main.views.sheetView[0].showGridLines = True

    # ── Sheet 2: Lookups (Reference Data) ──
    ws_ref = wb.create_sheet(title="Lookups")
    ws_ref.sheet_state = "hidden"
    ws_ref.views.sheetView[0].showGridLines = True

    # Headers for Sheet 1
    headers = [
        ("Full Name *", "full_name"),
        ("Gender", "gender"),
        ("Phone Number", "phone"),
        ("Email Address", "email"),
        ("Assigned Roles *", "roles"),
        ("Department", "department"),
        ("Subject", "subject"),
        ("Form Class", "form_class"),
        ("House Assigned", "house_assigned"),
        ("Password (Optional)", "password")
    ]

    header_font = Font(name="Segoe UI", size=10, bold=True, color="FFFFFF")
    header_fill = PatternFill(start_color="1E293B", end_color="1E293B", fill_type="solid")
    center_align = Alignment(horizontal="center", vertical="center", wrap_text=True)
    left_align = Alignment(horizontal="left", vertical="center")
    thin_border = Border(
        left=Side(style="thin", color="CBD5E1"),
        right=Side(style="thin", color="CBD5E1"),
        top=Side(style="thin", color="CBD5E1"),
        bottom=Side(style="thin", color="CBD5E1")
    )

    ws_main.row_dimensions[1].height = 28
    for col_idx, (display_name, key_name) in enumerate(headers, start=1):
        cell = ws_main.cell(row=1, column=col_idx, value=key_name)
        cell.font = header_font
        cell.fill = header_fill
        cell.alignment = center_align
        cell.border = thin_border

    # ── Populate Sheet 2 (Lookups) ──
    ref_columns = [
        ("Gender", lookups["genders"]),
        ("Roles", lookups["roles"]),
        ("Departments", lookups["departments"]),
        ("Subjects", lookups["subjects"]),
        ("Classes", lookups["classes"]),
        ("Houses", lookups["houses"])
    ]

    for col_idx, (col_name, items) in enumerate(ref_columns, start=1):
        ws_ref.cell(row=1, column=col_idx, value=col_name)
        for row_idx, item in enumerate(items, start=2):
            ws_ref.cell(row=row_idx, column=col_idx, value=item)

    # ── Configure Native Excel Data Validations on Sheet 1 (Rows 2 to 500) ──
    # Column B: Gender -> Lookups!$A$2:$A${len}
    if lookups["genders"]:
        max_r = len(lookups["genders"]) + 1
        dv_gender = DataValidation(type="list", formula1=f"=Lookups!$A$2:$A${max_r}", allow_blank=True)
        dv_gender.error = "Please select a valid gender from the dropdown."
        dv_gender.errorTitle = "Invalid Gender"
        dv_gender.prompt = "Select gender (Male or Female)"
        dv_gender.promptTitle = "Gender"
        ws_main.add_data_validation(dv_gender)
        dv_gender.add("B2:B500")

    # Column E: Roles -> Lookups!$B$2:$B${len}
    if lookups["roles"]:
        max_r = len(lookups["roles"]) + 1
        dv_roles = DataValidation(type="list", formula1=f"=Lookups!$B$2:$B${max_r}", allow_blank=True)
        dv_roles.prompt = "Select role(s) or type custom roles separated by |"
        dv_roles.promptTitle = "Assigned Roles"
        ws_main.add_data_validation(dv_roles)
        dv_roles.add("E2:E500")

    # Column F: Department -> Lookups!$C$2:$C${len}
    if lookups["departments"]:
        max_r = len(lookups["departments"]) + 1
        dv_dept = DataValidation(type="list", formula1=f"=Lookups!$C$2:$C${max_r}", allow_blank=True)
        dv_dept.prompt = "Select department from the school's registered departments"
        dv_dept.promptTitle = "School Department"
        ws_main.add_data_validation(dv_dept)
        dv_dept.add("F2:F500")

    # Column G: Subject -> Lookups!$D$2:$D${len}
    if lookups["subjects"]:
        max_r = len(lookups["subjects"]) + 1
        dv_subj = DataValidation(type="list", formula1=f"=Lookups!$D$2:$D${max_r}", allow_blank=True)
        dv_subj.prompt = "Select accredited/custom subject taught by this staff member"
        dv_subj.promptTitle = "Teaching Subject"
        ws_main.add_data_validation(dv_subj)
        dv_subj.add("G2:G500")

    # Column H: Form Class -> Lookups!$E$2:$E${len}
    if lookups["classes"]:
        max_r = len(lookups["classes"]) + 1
        dv_cls = DataValidation(type="list", formula1=f"=Lookups!$E$2:$E${max_r}", allow_blank=True)
        dv_cls.prompt = "Select class section for Form Master/Mistress assignment (e.g. 1ST3)"
        dv_cls.promptTitle = "Form Class"
        ws_main.add_data_validation(dv_cls)
        dv_cls.add("H2:H500")

    # Column I: House Assigned -> Lookups!$F$2:$F${len}
    if lookups["houses"]:
        max_r = len(lookups["houses"]) + 1
        dv_house = DataValidation(type="list", formula1=f"=Lookups!$F$2:$F${max_r}", allow_blank=True)
        dv_house.prompt = "Select boarding house for House Master/Mistress assignment"
        dv_house.promptTitle = "House Assigned"
        ws_main.add_data_validation(dv_house)
        dv_house.add("I2:I500")

    # Set generous column widths
    column_widths = {
        "A": 26,  # full_name
        "B": 14,  # gender
        "C": 18,  # phone
        "D": 28,  # email
        "E": 28,  # roles
        "F": 34,  # department
        "G": 32,  # subject
        "H": 16,  # form_class
        "I": 20,  # house_assigned
        "J": 22   # password
    }
    for col_letter, width in column_widths.items():
        ws_main.column_dimensions[col_letter].width = width

    output = io.BytesIO()
    wb.save(output)
    output.seek(0)

    clean_school = re.sub(r'[^a-zA-Z0-9_-]', '_', lookups["school_name"])
    filename = f"Staff_Onboarding_Template_{clean_school}.xlsx"

    return output.getvalue(), filename


async def parse_uploaded_staff_file(file: UploadFile, max_bytes: int = 15 * 1024 * 1024) -> Tuple[List[Dict[str, str]], str]:
    """
    Seamlessly parses both .xlsx and .csv files.
    Returns a list of cleaned row dictionaries with normalized keys and safe filename.
    """
    if not file or not file.filename:
        raise HTTPException(status_code=400, detail="No file uploaded or missing filename.")

    filename_lower = file.filename.lower().strip()
    safe_name = sanitize_filename(file.filename, default="staff_upload.xlsx")
    content = await file.read()

    if len(content) > max_bytes:
        raise HTTPException(status_code=400, detail=f"File exceeds maximum allowed limit ({max_bytes // (1024 * 1024)}MB).")

    rows: List[Dict[str, str]] = []

    # ── 1. Process Excel (.xlsx) ──
    if filename_lower.endswith(".xlsx"):
        try:
            wb = openpyxl.load_workbook(io.BytesIO(content), data_only=True)
            # Pick first visible worksheet
            ws = wb.active
            for s in wb.worksheets:
                if s.sheet_state != "hidden":
                    ws = s
                    break

            raw_rows = list(ws.iter_rows(values_only=True))
            if not raw_rows or len(raw_rows) < 2:
                return [], safe_name

            header_row = [str(cell).strip() if cell is not None else "" for cell in raw_rows[0]]
            for row in raw_rows[1:]:
                # Skip completely empty rows
                if not any(row):
                    continue
                row_dict = {}
                for idx, h in enumerate(header_row):
                    if not h:
                        continue
                    val = row[idx] if idx < len(row) else ""
                    row_dict[str(h).strip().lower()] = str(val).strip() if val is not None else ""
                if any(row_dict.values()):
                    rows.append(row_dict)

            return rows, safe_name
        except Exception as e:
            raise HTTPException(status_code=400, detail=f"Error reading Excel workbook: {str(e)}")

    # ── 2. Process CSV (.csv) ──
    elif filename_lower.endswith(".csv"):
        import csv
        decoded = decode_csv_bytes(content, max_bytes=max_bytes)
        stream = io.StringIO(decoded)
        reader = csv.DictReader(stream)
        for r in reader:
            clean_r = {str(k).strip().lower(): str(v).strip() if v is not None else "" for k, v in r.items() if k}
            if any(clean_r.values()):
                rows.append(clean_r)
        return rows, safe_name

    else:
        raise HTTPException(
            status_code=400,
            detail="Unsupported file format. Please upload an Excel workbook (.xlsx) or CSV (.csv)."
        )
