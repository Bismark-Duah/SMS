"""
broadsheet_export_service.py - Enterprise Academic Broadsheet & Examination Scores Export Engine.
Provides filtered, preset-driven exports for:
1. Master Class Scores Broadsheet (All subjects in matrix grid)
2. Continuous Assessment (SBA) / WAEC Assessment Roster
3. Subject Performance & League Table

Supports both styled .xlsx (openpyxl) and universal .csv with strict school-level isolation.
100% Offline-First.
"""

import io
import csv
import re
from datetime import datetime
from typing import Any, Dict, List, Optional, Tuple
from sqlalchemy.orm import Session, joinedload
from sqlalchemy import func, desc

import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter

from ..models import Student, Score, Subject, ClassSection, Semester, School
from .import_export_service import sanitize_csv_cell


def get_broadsheet_lookups(db: Session, school_id: Optional[int]) -> Dict[str, Any]:
    """Retrieves classes and active semesters for broadsheet filtering."""
    cls_query = db.query(ClassSection)
    if school_id is not None:
        cls_query = cls_query.filter(ClassSection.school_id == school_id)
    classes = cls_query.order_by(ClassSection.name.asc()).all()

    sem_query = db.query(Semester).options(joinedload(Semester.academic_year))
    semesters = sem_query.order_by(desc(Semester.is_current), desc(Semester.id)).all()

    sub_query = db.query(Subject)
    if school_id is not None:
        sub_query = sub_query.filter(
            (Subject.school_id == school_id) | (Subject.school_id.is_(None))
        )
    subjects = sub_query.order_by(Subject.name.asc()).all()

    return {
        "classes": [{"id": c.id, "name": c.name} for c in classes],
        "semesters": [{"id": s.id, "name": f"{s.academic_year.label if s.academic_year else ''} - {s.name or ''}", "is_current": bool(s.is_current)} for s in semesters],
        "subjects": [{"id": sub.id, "name": sub.name, "code": sub.code} for sub in subjects]
    }


def generate_broadsheet_export_dataset(
    db: Session,
    school_id: Optional[int],
    preset_key: str = "master_broadsheet",
    file_format: str = "xlsx",
    filters: Optional[Dict[str, Any]] = None
) -> Tuple[bytes, str, str]:
    """
    Generates preset-driven academic score exports.
    """
    filters = filters or {}
    school = db.query(School).filter(School.id == school_id).first() if school_id else None
    school_slug = re.sub(r"[^a-zA-Z0-9_-]", "_", school.name if school else "Institutional")
    timestamp = datetime.now().strftime("%Y%m%d_%H%M")

    # 1. Resolve Semester / Term
    sem_id = filters.get("semester_id")
    if not sem_id or str(sem_id).strip() == "":
        cur_sem = db.query(Semester).filter(Semester.is_current == True).first()
        if not cur_sem:
            cur_sem = db.query(Semester).order_by(desc(Semester.id)).first()
        sem_id = cur_sem.id if cur_sem else None
    else:
        try:
            sem_id = int(sem_id)
        except (ValueError, TypeError):
            sem_id = None

    class_id = filters.get("class_id")
    subject_id = filters.get("subject_id")

    # ── Preset 1: SBA / WAEC Assessment Roster ──
    if preset_key == "sba_waec":
        title = "SBA_WAEC_Assessment_Roster"
        columns = [
            ("Student Code", "student_code"),
            ("BECE Index No", "bece_index_number"),
            ("Full Name", "full_name"),
            ("Gender", "gender"),
            ("Class Section", "class_name"),
            ("Subject", "subject_name"),
            ("Class Score (30%)", "class_score"),
            ("Exam Score (70%)", "exam_score"),
            ("Total Score (100%)", "total_score"),
            ("Grade", "grade"),
            ("Remark", "remark")
        ]

        query = db.query(Score).join(Student, Score.student_id == Student.id)\
                               .join(Subject, Score.subject_id == Subject.id)\
                               .options(
                                   joinedload(Score.student).joinedload(Student.class_section),
                                   joinedload(Score.subject)
                               )
        if school_id is not None:
            query = query.filter(Student.school_id == school_id)
        if sem_id:
            query = query.filter(Score.semester_id == sem_id)
        if class_id:
            try:
                query = query.filter(Student.class_section_id == int(class_id))
            except (ValueError, TypeError):
                pass
        if subject_id:
            try:
                query = query.filter(Score.subject_id == int(subject_id))
            except (ValueError, TypeError):
                pass

        scores = query.order_by(Student.class_section_id.asc(), Student.full_name.asc()).all()
        rows_data = []
        for sc in scores:
            st = sc.student
            sub = sc.subject
            rows_data.append({
                "student_code": st.student_code if st else "",
                "bece_index_number": st.bece_index_number if st else "",
                "full_name": st.full_name if st else "",
                "gender": st.gender if st else "",
                "class_name": st.class_section.name if (st and st.class_section) else "",
                "subject_name": sub.name if sub else "",
                "class_score": round(sc.class_score or 0.0, 1),
                "exam_score": round(sc.exam_score or 0.0, 1),
                "total_score": round(sc.total_score or 0.0, 1),
                "grade": sc.grade or "",
                "remark": sc.remark or ""
            })

        return _render_simple_dataset(rows_data, columns, f"{title}_{school_slug}_{timestamp}", file_format)

    # ── Preset 2: Subject Performance & League Table ──
    elif preset_key == "subject_summary":
        title = "Subject_Performance_League_Table"
        columns = [
            ("Subject", "subject_name"),
            ("Class Section", "class_name"),
            ("Candidates Assessed", "candidates"),
            ("Pass Count (>=50)", "pass_count"),
            ("Fail Count (<50)", "fail_count"),
            ("Pass Rate (%)", "pass_rate"),
            ("Highest Mark", "highest"),
            ("Lowest Mark", "lowest"),
            ("Class Average", "average")
        ]

        query = db.query(Score).join(Student, Score.student_id == Student.id)\
                               .join(Subject, Score.subject_id == Subject.id)\
                               .options(
                                   joinedload(Score.student).joinedload(Student.class_section),
                                   joinedload(Score.subject)
                               )
        if school_id is not None:
            query = query.filter(Student.school_id == school_id)
        if sem_id:
            query = query.filter(Score.semester_id == sem_id)
        if class_id:
            try:
                query = query.filter(Student.class_section_id == int(class_id))
            except (ValueError, TypeError):
                pass

        scores = query.all()

        # Group by (subject_id, class_section_id)
        grouped: Dict[Tuple[int, Optional[int]], List[Score]] = {}
        for sc in scores:
            cls_id = sc.student.class_section_id if sc.student else None
            key = (sc.subject_id, cls_id)
            grouped.setdefault(key, []).append(sc)

        rows_data = []
        for (sub_k, cls_k), sc_list in grouped.items():
            if not sc_list:
                continue
            sub_name = sc_list[0].subject.name if sc_list[0].subject else f"Subject #{sub_k}"
            cls_name = sc_list[0].student.class_section.name if (sc_list[0].student and sc_list[0].student.class_section) else "All Classes"

            totals = [s.total_score for s in sc_list if s.total_score is not None]
            cands = len(totals)
            if cands == 0:
                continue

            passes = sum(1 for t in totals if t >= 50.0)
            fails = cands - passes
            rate = round((passes / cands) * 100.0, 1)
            high = max(totals)
            low = min(totals)
            avg = round(sum(totals) / cands, 1)

            rows_data.append({
                "subject_name": sub_name,
                "class_name": cls_name,
                "candidates": cands,
                "pass_count": passes,
                "fail_count": fails,
                "pass_rate": f"{rate}%",
                "highest": high,
                "lowest": low,
                "average": avg
            })

        rows_data.sort(key=lambda x: (x["class_name"], x["subject_name"]))
        return _render_simple_dataset(rows_data, columns, f"{title}_{school_slug}_{timestamp}", file_format)

    # ── Preset 3: Master Class Scores Broadsheet ──
    else:
        title = "Master_Class_Scores_Broadsheet"
        # 1. Fetch Students
        st_query = db.query(Student).filter(Student.is_active == True)
        if school_id is not None:
            st_query = st_query.filter(Student.school_id == school_id)
        if class_id:
            try:
                st_query = st_query.filter(Student.class_section_id == int(class_id))
            except (ValueError, TypeError):
                pass

        students = st_query.options(joinedload(Student.class_section)).order_by(Student.full_name.asc()).all()
        if not students:
            # Empty fallback
            return _render_simple_dataset([], [("Notice", "notice")], f"{title}_{school_slug}_{timestamp}", file_format)

        st_ids = [s.id for s in students]

        # 2. Fetch Scores for these students in selected semester
        sc_query = db.query(Score).filter(Score.student_id.in_(st_ids)).options(joinedload(Score.subject))
        if sem_id:
            sc_query = sc_query.filter(Score.semester_id == sem_id)
        all_scores = sc_query.all()

        # Group scores by student and subject
        student_scores: Dict[int, Dict[int, Score]] = {}
        subject_map: Dict[int, str] = {}
        for sc in all_scores:
            student_scores.setdefault(sc.student_id, {})[sc.subject_id] = sc
            if sc.subject and sc.subject.name:
                subject_map[sc.subject_id] = sc.subject.name.strip()

        # Sort subjects alphabetically
        sorted_subject_ids = sorted(subject_map.keys(), key=lambda sid: subject_map[sid])

        # Build dynamic columns
        # Core info
        columns = [
            ("Student Code", "student_code"),
            ("Full Name", "full_name"),
            ("Gender", "gender"),
            ("Class Section", "class_name")
        ]
        # Subject score columns
        for sid in sorted_subject_ids:
            sname = subject_map[sid]
            columns.append((f"{sname} (Total)", f"sub_{sid}_score"))
            columns.append((f"{sname} (Grade)", f"sub_{sid}_grade"))

        columns.append(("Grand Total", "grand_total"))
        columns.append(("Average (%)", "average_score"))
        columns.append(("Class Rank", "class_rank"))

        # Build student rows
        student_summaries = []
        for s in students:
            s_scores = student_scores.get(s.id, {})
            tot = 0.0
            count = 0
            row_dict = {
                "student_code": s.student_code or "",
                "full_name": s.full_name or "",
                "gender": s.gender or "",
                "class_name": s.class_section.name if s.class_section else ""
            }

            for sid in sorted_subject_ids:
                sc = s_scores.get(sid)
                if sc and sc.total_score is not None:
                    row_dict[f"sub_{sid}_score"] = round(sc.total_score, 1)
                    row_dict[f"sub_{sid}_grade"] = sc.grade or ""
                    tot += sc.total_score
                    count += 1
                else:
                    row_dict[f"sub_{sid}_score"] = ""
                    row_dict[f"sub_{sid}_grade"] = ""

            avg = round(tot / count, 1) if count > 0 else 0.0
            row_dict["grand_total"] = round(tot, 1)
            row_dict["average_score"] = avg
            student_summaries.append(row_dict)

        # Calculate ranks by average score descending
        student_summaries.sort(key=lambda x: x["average_score"], reverse=True)
        for rank_idx, item in enumerate(student_summaries, start=1):
            item["class_rank"] = f"#{rank_idx}"

        # Re-sort by name or class if desired
        student_summaries.sort(key=lambda x: (x["class_name"], x["full_name"]))

        return _render_simple_dataset(student_summaries, columns, f"{title}_{school_slug}_{timestamp}", file_format)


def _render_simple_dataset(
    rows_data: List[Dict[str, Any]],
    columns: List[Tuple[str, str]],
    filename_base: str,
    file_format: str
) -> Tuple[bytes, str, str]:
    """Helper to render styled Excel or CSV for broadsheets."""
    if file_format.lower() == "xlsx":
        wb = openpyxl.Workbook()
        ws = wb.active
        ws.title = "Broadsheet"
        ws.views.sheetView[0].showGridLines = True

        header_font = Font(name="Segoe UI", size=10, bold=True, color="FFFFFF")
        header_fill = PatternFill(start_color="1E293B", end_color="1E293B", fill_type="solid")
        cell_font = Font(name="Segoe UI", size=10)
        center_align = Alignment(horizontal="center", vertical="center")
        left_align = Alignment(horizontal="left", vertical="center")
        right_align = Alignment(horizontal="right", vertical="center")
        thin_border = Border(
            left=Side(style="thin", color="CBD5E1"),
            right=Side(style="thin", color="CBD5E1"),
            top=Side(style="thin", color="CBD5E1"),
            bottom=Side(style="thin", color="CBD5E1")
        )
        zebra_fill = PatternFill(start_color="F8FAFC", end_color="F8FAFC", fill_type="solid")

        # Header
        ws.row_dimensions[1].height = 26
        for col_idx, (col_label, col_key) in enumerate(columns, start=1):
            cell = ws.cell(row=1, column=col_idx, value=col_label)
            cell.font = header_font
            cell.fill = header_fill
            cell.alignment = center_align
            cell.border = thin_border

        # Data rows
        for row_idx, r in enumerate(rows_data, start=2):
            ws.row_dimensions[row_idx].height = 20
            is_even = (row_idx % 2 == 0)

            for col_idx, (col_label, col_key) in enumerate(columns, start=1):
                val = r.get(col_key, "")
                cell = ws.cell(row=row_idx, column=col_idx, value=val)
                cell.font = cell_font
                cell.border = thin_border
                if is_even:
                    cell.fill = zebra_fill

                if isinstance(val, (int, float)):
                    cell.alignment = right_align
                elif col_key in ("student_code", "gender", "class_rank", "grade", "pass_rate") or col_key.endswith("_grade"):
                    cell.alignment = center_align
                else:
                    cell.alignment = left_align

        # Auto-fit columns
        for col in ws.columns:
            max_len = 0
            col_letter = get_column_letter(col[0].column)
            for cell in col:
                val_str = str(cell.value or "")
                if len(val_str) > max_len:
                    max_len = len(val_str)
            ws.column_dimensions[col_letter].width = max(max_len + 4, 11)

        output = io.BytesIO()
        wb.save(output)
        output.seek(0)
        return output.getvalue(), f"{filename_base}.xlsx", "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"

    else:
        output = io.StringIO()
        writer = csv.writer(output, quoting=csv.QUOTE_MINIMAL)
        writer.writerow([c[0] for c in columns])

        for r in rows_data:
            row_vals = [sanitize_csv_cell(r.get(c[1], "")) for c in columns]
            writer.writerow(row_vals)

        csv_bytes = output.getvalue().encode("utf-8-sig")
        return csv_bytes, f"{filename_base}.csv", "text/csv; charset=utf-8"
